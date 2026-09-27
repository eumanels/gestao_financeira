"""Conversão entre o arquivo Excel e DataFrames do pandas.

Estrutura: UMA ABA POR USUÁRIO. Cada aba é uma tabela única ("formato longo") em que a
coluna ``Registro`` diz o que cada linha representa:

    SALARIO    -> salário de um mês (Mes_Referencia + Valor). Uma linha por mês.
    COFRINHO   -> movimentação da reserva (Tipo = Depósito/Retirada, Valor sempre positivo).
    DIVIDA     -> cadastro de uma dívida/despesa (recorrente ou parcelada).
    PAGAMENTO  -> histórico: uma parcela paga de uma dívida (ID_Divida) em um mês.

Por que uma tabela só? Porque a planilha base tem apenas a linha de cabeçalhos: com um
cabeçalho único por aba, a aplicação consegue criar/validar a estrutura automaticamente e
o arquivo continua fácil de filtrar no Excel (filtro na coluna "Registro").

Abas cujo nome começa com "_" (ex.: "_modelo") são ignoradas na lista de usuários.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Cabeçalhos oficiais (ordem em que são gravados na planilha)
COLUMNS: list[str] = [
    "Registro",
    "ID",
    "Mes_Referencia",
    "Data",
    "Descricao",
    "Categoria",
    "Tipo",
    "Valor",
    "Parcelas_Total",
    "Parcelas_Pagas",
    "Status_Mes",
    "Ativa",
    "ID_Divida",
]
TEXT_COLUMNS = [
    "Registro", "ID", "Mes_Referencia", "Descricao", "Categoria",
    "Tipo", "Status_Mes", "Ativa", "ID_Divida",
]
NUMERIC_COLUMNS = ["Valor", "Parcelas_Total", "Parcelas_Pagas"]

# Valores da coluna Registro
REG_SALARIO = "SALARIO"
REG_COFRINHO = "COFRINHO"
REG_DIVIDA = "DIVIDA"
REG_PAGAMENTO = "PAGAMENTO"

TEMPLATE_SHEET = "_modelo"
# Nomes padrão do Excel: se a aba estiver vazia, é tratada como modelo e não como usuário
DEFAULT_SHEET_NAMES = {"planilha1", "plan1", "sheet1", "folha1", "modelo"}

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_MONEY_FORMAT = '"R$" #,##0.00'
_DATE_FORMAT = "dd/mm/yyyy hh:mm"
_COLUMN_WIDTHS = {
    "Registro": 12, "ID": 11, "Mes_Referencia": 15, "Data": 17, "Descricao": 28,
    "Categoria": 16, "Tipo": 12, "Valor": 14, "Parcelas_Total": 14, "Parcelas_Pagas": 14,
    "Status_Mes": 12, "Ativa": 8, "ID_Divida": 11,
}


class SchemaError(Exception):
    """A aba existe mas não segue a estrutura esperada."""


# --------------------------------------------------------------------------- #
# DataFrame vazio / normalização de tipos
# --------------------------------------------------------------------------- #
def empty_frame() -> pd.DataFrame:
    """DataFrame vazio já com todas as colunas e tipos corretos."""
    return normalize(pd.DataFrame(columns=COLUMNS))


def _to_month(value) -> str:
    """Converte qualquer representação de mês para o texto 'AAAA-MM'."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return f"{value.year:04d}-{value.month:02d}"
    text = str(value).strip()
    match = re.match(r"^(\d{4})[-/](\d{1,2})", text)
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"
    match = re.match(r"^(\d{1,2})[-/](\d{4})$", text)  # formato MM/AAAA digitado à mão
    if match:
        return f"{int(match.group(2)):04d}-{int(match.group(1)):02d}"
    return text


def normalize(df: pd.DataFrame) -> pd.DataFrame:
    """Garante colunas, ordem e tipos consistentes (texto, número e data)."""
    df = df.copy()
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = None
    df = df[COLUMNS]
    for col in TEXT_COLUMNS:
        df[col] = df[col].map(
            lambda v: "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
        ).astype(object)
    df["Mes_Referencia"] = df["Mes_Referencia"].map(_to_month).astype(object)
    df["Registro"] = df["Registro"].str.upper()
    df["Ativa"] = df["Ativa"].str.upper()
    for col in NUMERIC_COLUMNS:
        df[col] = pd.to_numeric(df[col], errors="coerce").astype(float)
    df["Data"] = pd.to_datetime(df["Data"], errors="coerce", dayfirst=True)
    # Linhas sem tipo de registro não significam nada para a aplicação
    df = df[df["Registro"] != ""]
    return df.reset_index(drop=True)


# --------------------------------------------------------------------------- #
# Leitura
# --------------------------------------------------------------------------- #
def open_workbook(data: bytes | None) -> Workbook:
    """Abre o .xlsx a partir dos bytes (ou cria um novo, com a aba modelo)."""
    if not data:
        wb = Workbook()
        ws = wb.active
        ws.title = TEMPLATE_SHEET
        _write_rows(ws, empty_frame())
        return wb
    try:
        return load_workbook(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - openpyxl lança vários tipos
        raise SchemaError(
            f"O arquivo não é uma planilha .xlsx válida ou está corrompido ({exc})."
        ) from exc


def sheet_to_frame(ws: Worksheet) -> pd.DataFrame:
    """Lê uma aba e devolve um DataFrame normalizado.

    * Aba totalmente vazia ou só com cabeçalhos -> DataFrame vazio (estrutura criada
      automaticamente no primeiro salvamento).
    * Aba com dados mas sem as colunas obrigatórias -> ``SchemaError``.
    """
    rows = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    if not rows:
        return empty_frame()

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    lookup = {name.lower(): name for name in COLUMNS}
    canonical = [lookup.get(h.lower(), h) for h in header]
    data_rows = rows[1:]

    missing = [c for c in COLUMNS if c not in canonical]
    if missing and data_rows:
        raise SchemaError(
            f"A aba '{ws.title}' tem dados, mas faltam as colunas: {', '.join(missing)}. "
            f"Cabeçalhos esperados: {', '.join(COLUMNS)}."
        )
    if not data_rows:
        return empty_frame()

    df = pd.DataFrame(data_rows, columns=canonical)
    df = df.loc[:, ~df.columns.duplicated()]
    return normalize(df)


def is_template_sheet(ws: Worksheet) -> bool:
    """Abas que NÃO representam usuários."""
    if ws.title.startswith("_"):
        return True
    if ws.title.strip().lower() in DEFAULT_SHEET_NAMES:
        return ws.max_row <= 1  # aba padrão vazia/só cabeçalho = modelo
    return False


def load_user_frames(data: bytes | None) -> dict[str, pd.DataFrame]:
    """Lê todas as abas de usuários do arquivo -> ``{nome_usuario: DataFrame}``."""
    wb = open_workbook(data)
    return {ws.title: sheet_to_frame(ws) for ws in wb.worksheets if not is_template_sheet(ws)}


# --------------------------------------------------------------------------- #
# Escrita
# --------------------------------------------------------------------------- #
def _cell_value(value):
    """Converte valores do pandas/numpy em tipos que o openpyxl entende."""
    if value is None:
        return None
    if isinstance(value, pd.Timestamp):
        return None if pd.isna(value) else value.to_pydatetime()
    if isinstance(value, float) and pd.isna(value):
        return None
    if hasattr(value, "item"):  # numpy.int64 / numpy.float64 -> int/float do Python
        value = value.item()
    if isinstance(value, str) and value == "":
        return None  # célula vazia no Excel
    return value


def _write_rows(ws: Worksheet, df: pd.DataFrame) -> None:
    """Escreve cabeçalho + linhas na aba, com formatação básica."""
    ws.append(COLUMNS)
    for row in df[COLUMNS].itertuples(index=False):
        ws.append([_cell_value(v) for v in row])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for idx, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = _COLUMN_WIDTHS.get(name, 14)

    valor_col = COLUMNS.index("Valor") + 1
    data_col = COLUMNS.index("Data") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=valor_col).number_format = _MONEY_FORMAT
        ws.cell(row=r, column=data_col).number_format = _DATE_FORMAT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_user_sheet(data: bytes | None, user: str, df: pd.DataFrame) -> bytes:
    """Substitui SOMENTE a aba do usuário, preservando as demais abas do arquivo."""
    wb = open_workbook(data)
    if user in wb.sheetnames:
        ws = wb[user]
        ws.delete_rows(1, ws.max_row)  # limpa a aba mantendo sua posição no arquivo
    else:
        ws = wb.create_sheet(title=user)
    _write_rows(ws, normalize(df))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Validação de nome de usuário (= nome da aba)
# --------------------------------------------------------------------------- #
def validate_user_name(name: str, existing: list[str]) -> str:
    """Valida o nome conforme as regras de nomes de aba do Excel."""
    name = (name or "").strip()
    if not name:
        raise ValueError("Informe um nome de usuário.")
    if len(name) > 31:
        raise ValueError("O nome pode ter no máximo 31 caracteres (limite do Excel).")
    if _INVALID_SHEET_CHARS.search(name):
        raise ValueError("O nome não pode conter os caracteres  [ ] : * ? / \\")
    if name.startswith("_") or name.startswith("'"):
        raise ValueError("O nome não pode começar com _ ou apóstrofo.")
    if name.lower() in DEFAULT_SHEET_NAMES:
        raise ValueError("Esse nome é reservado. Escolha outro.")
    duplicate = next((e for e in existing if e.lower() == name.lower()), None)
    if duplicate:  # o Excel não diferencia maiúsculas/minúsculas em nomes de aba
        raise ValueError(f"Já existe um usuário chamado '{duplicate}'.")
    return name
