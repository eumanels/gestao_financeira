"""Testes das regras de negócio e da leitura/escrita da planilha.

Rodar:  pip install -r requirements-dev.txt  &&  pytest -q
"""

from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from finance import operations as ops
from finance import repository as repo
from finance.calculations import brl, debts_for_month, month_summary, piggy_balance
from finance.storage import ConflictError, LocalBackend

M1, M2 = "2026-09", "2026-10"


def test_base_sheet_with_only_headers_is_empty_user():
    wb = Workbook()
    wb.active.title = "Ana"
    wb.active.append(repo.COLUMNS)
    buf = io.BytesIO()
    wb.save(buf)
    frames = repo.load_user_frames(buf.getvalue())
    assert list(frames) == ["Ana"]
    assert frames["Ana"].empty and list(frames["Ana"].columns) == repo.COLUMNS


def test_default_empty_sheet_is_hidden_and_new_file_is_created():
    wb = Workbook()  # aba "Sheet" vazia
    wb.active.title = "Planilha1"
    buf = io.BytesIO()
    wb.save(buf)
    assert repo.load_user_frames(buf.getvalue()) == {}
    assert repo.load_user_frames(None) == {}


def test_sheet_with_data_and_wrong_headers_raises():
    wb = Workbook()
    wb.active.title = "Bob"
    wb.active.append(["Nome", "Valor"])
    wb.active.append(["x", 1])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(repo.SchemaError):
        repo.load_user_frames(buf.getvalue())


def test_full_flow_and_balance():
    df = repo.empty_frame()
    df = ops.set_salary(df, M1, 5000)
    df = ops.set_salary(df, M1, 5200)  # atualiza, não duplica
    df = ops.add_debt(df, "Aluguel", ops.TIPO_RECORRENTE, 1500, M1, category="Aluguel")
    df = ops.add_debt(df, "Notebook", ops.TIPO_PARCELADA, 400, M1, total_installments=3, paid_installments=1)
    rent = df.loc[df["Descricao"] == "Aluguel", "ID"].iloc[0]
    note = df.loc[df["Descricao"] == "Notebook", "ID"].iloc[0]

    s = month_summary(df, M1)
    assert (s.salary, s.total_debts, s.paid, s.balance) == (5200, 1900, 0, 5200)

    df = ops.pay_debt(df, rent, M1)
    df = ops.pay_debt(df, note, M1)
    with pytest.raises(ops.ValidationError):
        ops.pay_debt(df, rent, M1)  # não paga duas vezes no mesmo mês
    s = month_summary(df, M1)
    assert (s.paid, s.balance, s.pending) == (1900, 3300, 0)

    # Última parcela do notebook no mês seguinte -> quitada
    df = ops.pay_debt(df, note, M2)
    debt = df[(df["Registro"] == repo.REG_DIVIDA) & (df["ID"] == note)].iloc[0]
    assert debt["Parcelas_Pagas"] == 3 and debt["Ativa"] == ops.NAO
    # Mês 3: notebook não aparece mais; aluguel continua
    assert list(debts_for_month(df, "2026-11")["Descricao"]) == ["Aluguel"]

    # Desfazer reativa a parcelada
    df = ops.undo_payment(df, note, M2)
    debt = df[(df["Registro"] == repo.REG_DIVIDA) & (df["ID"] == note)].iloc[0]
    assert debt["Parcelas_Pagas"] == 2 and debt["Ativa"] == ops.SIM


def test_piggy_bank():
    df = ops.move_piggy_bank(repo.empty_frame(), ops.TIPO_DEPOSITO, 300, M1)
    df = ops.move_piggy_bank(df, ops.TIPO_RETIRADA, 100, M1, current_balance=piggy_balance(df))
    assert piggy_balance(df) == 200
    with pytest.raises(ops.ValidationError):
        ops.move_piggy_bank(df, ops.TIPO_RETIRADA, 500, M1, current_balance=piggy_balance(df))


def test_write_preserves_other_sheets_and_roundtrip(tmp_path):
    df = ops.set_salary(repo.empty_frame(), M1, 1000)
    df = ops.add_debt(df, "Cartão", ops.TIPO_PARCELADA, 250.5, M1, total_installments=10)
    data = repo.write_user_sheet(None, "Ana", df)
    data = repo.write_user_sheet(data, "Bob", ops.set_salary(repo.empty_frame(), M1, 3000))
    data = repo.write_user_sheet(data, "Ana", ops.set_salary(df, M1, 1100))  # reescreve só Ana

    frames = repo.load_user_frames(data)
    assert set(frames) == {"Ana", "Bob"}
    assert month_summary(frames["Ana"], M1).salary == 1100
    assert month_summary(frames["Bob"], M1).salary == 3000
    assert frames["Ana"].loc[frames["Ana"]["Registro"] == "DIVIDA", "Parcelas_Total"].iloc[0] == 10
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == repo.TEMPLATE_SHEET

    backend = LocalBackend(str(tmp_path / "sub" / "f.xlsx"))
    assert backend.read() == (None, None)
    v1 = backend.write(data, None)
    content, v = backend.read()
    assert content == data and v == v1
    with pytest.raises(ConflictError):
        backend.write(data, "versao-antiga")


def test_user_name_validation():
    assert repo.validate_user_name(" Ana ", []) == "Ana"
    for bad in ["", "a/b", "x" * 32, "_x", "Planilha1"]:
        with pytest.raises(ValueError):
            repo.validate_user_name(bad, [])
    with pytest.raises(ValueError):
        repo.validate_user_name("ana", ["Ana"])


def test_brl():
    assert brl(1234.5) == "R$ 1.234,50"
    assert brl(-10) == "-R$ 10,00"
    assert brl(None) == "—"
